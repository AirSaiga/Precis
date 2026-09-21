# SPDX-License-Identifier: Apache-2.0
#
# Copyright 2026 Precis Team
#
# Licensed under the Apache License, Version 2.0 (the "License");
# you may not use this file except in compliance with the License.
# You may obtain a copy of the License at
#
#     http://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing, software
# distributed under the License is distributed on an "AS IS" BASIS,
# WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
# See the License for the specific language governing permissions and
# limitations under the License.
"""PyPI 发版视角审计（2026-09-21）发现的复现验证测试。

历史设计：每条用例断言【期望的正确行为】并标 ``xfail(strict=True)``——
xfail = 缺陷坐实；修复后 XPASS 转红提醒摘标。

2026-09-21 修复批次（R1-R8/H10 + R2 收窄）已全部落地：本文件的 xfail
标记已全部摘除，11 用例转为正式回归测试，长期守卫对应契约。

编号对应审计报告：R1-R8 为红线级，H9/H10 为高置信。R8（分块坏行行号
丢失）与 H9（MCP isError）不在此文件——前者的回归在
test_chunked_loader_deep.py（透传断言），后者在 tests/integration/
test_mcp_server.py（in-process isError 断言）。
"""

from __future__ import annotations

import io
import os
import subprocess
import sys
from pathlib import Path

import pytest

# 子进程用例须测仓库源码而非 site-packages 里可能陈旧的安装副本：
# 显式把 backend 源码根注入 PYTHONPATH（本文件位于 backend/tests/unit/ 下）
_BACKEND_SRC_ROOT = str(Path(__file__).resolve().parents[2])

# ---------------------------------------------------------------------------
# R3（红线#3）：validate 单发、无 --manifest、未打开项目 → 退出码应为 2（参数/使用错误）
# 当前坐实：exit 1（语义"发现数据违规"），CI 按 0/1/2 分流会误判。
# ---------------------------------------------------------------------------


def test_r3_validate_without_manifest_exits_2(tmp_path: Path) -> None:
    proc = subprocess.run(
        [sys.executable, "-m", "app.cli.shell.main", "validate"],
        cwd=tmp_path,  # 干净目录：无活跃项目、无清单
        capture_output=True,
        text=True,
        encoding="utf-8",
        errors="replace",
        timeout=120,
        env={**os.environ, "PYTHONPATH": _BACKEND_SRC_ROOT},
    )
    assert proc.returncode == 2, f"使用错误应退出码 2，实际 {proc.returncode}；stderr: {proc.stderr[:200]}"


# ---------------------------------------------------------------------------
# R4（红线#4）：config set 的类型推断不应接受 inf/nan/下划线整数
# 当前坐实：float('inf')/float('nan')/int('1_0') 全部通过并写入 YAML，
# timeout_seconds=.inf 使 validate 读回 int(inf) 抛 OverflowError（项目从此无法校验）。
# ---------------------------------------------------------------------------


def test_r4_parse_rejects_inf() -> None:
    from app.cli.shared_services.config_ops import parse_config_value

    ok, value, _ = parse_config_value("inf")
    assert not ok or value != float("inf"), "inf 不应被接受为可写配置值"


def test_r4_parse_rejects_nan() -> None:
    from app.cli.shared_services.config_ops import parse_config_value

    ok, value, _ = parse_config_value("nan")
    assert not ok or value == value, "nan 不应被接受为可写配置值（NaN != NaN）"


def test_r4_parse_rejects_underscore_int() -> None:
    from app.cli.shared_services.config_ops import parse_config_value

    ok, value, _ = parse_config_value("1_0")
    assert not (ok and value == 10), "'1_0' 不应被静默转换为整数 10"


# ---------------------------------------------------------------------------
# R5（红线#5）：校验结果渲染须转义数据值/消息中的 rich 标记
# 当前坐实：formatter.format_validation_result 拼接未经 markup_escape——
# 数据值 '[none]' 被当样式标签吞掉；消息含 '[/x]' 时 MarkupError 使
# "发现违规"崩成 exit 2。
# ---------------------------------------------------------------------------


def _render(payload_lines: str) -> str:
    from rich.console import Console

    buf = io.StringIO()
    Console(file=buf, width=200, force_terminal=False, legacy_windows=False).print(payload_lines)
    return buf.getvalue()


def test_r5_mismatched_close_tag_does_not_crash() -> None:
    from app.cli.shell.formatter import Formatter

    lines = Formatter.format_validation_result(
        [
            {
                "error_type": "RegexViolation",
                "table": "orders",
                "column": "code",
                "row_index": 0,
                "cell_value": "X[/x]Y",
                "error_message": "值 'X[/x]Y' 不符合正则表达式模式",
            }
        ]
    )
    rendered = _render(lines)  # 当前在此抛 rich.errors.MarkupError
    assert "X[/x]Y" in rendered or "X" in rendered


def test_r5_bracket_value_not_swallowed() -> None:
    from app.cli.shell.formatter import Formatter

    lines = Formatter.format_validation_result(
        [
            {
                "error_type": "AllowedValuesViolation",
                "table": "orders",
                "column": "status",
                "row_index": 0,
                "cell_value": "[none]",
                "error_message": "值 '[none]' 不在允许值集合内",
            }
        ]
    )
    assert "[none]" in _render(lines), "字面 [none] 应原样呈现给用户"


# ---------------------------------------------------------------------------
# R2（红线#2，经 strict 复核后收窄）：infer_schema 的 date 分支采用 fromisoformat
# 宽松集（Python>=3.11 接受 ISO 周日期），而下游 DateType 严格 strptime('%Y-%m-%d')。
# 注：YYYYMMDD 主场景不触发（integer 分支先命中，代理报告此处为误报）；
# 触发形态收窄为 ISO 周日期 '2025-W01-1' 等——被推成 date 后下游整列误报。
# ---------------------------------------------------------------------------


def test_r2_infer_rejects_week_date_text() -> None:
    from app.shared.services.schema_inference import infer_column_type

    assert infer_column_type(["2025-W01-1", "2025-W02-1", "2025-W03-1"]) != "date", (
        "ISO 周日期与下游 %Y-%m-%d 严格格式不一致，不应推断为 date"
    )


# ---------------------------------------------------------------------------
# R1（红线#1）：分块 Excel 流式解析器的合并区域首行值应保留日期语义
# 当前坐实：不读 styles.xml 数字格式，日期单元格返回序列号 int（45306），
# 跨块合并填充后同列混合 Timestamp 与 int → 日期约束误报、两路径判定不一致。
# ---------------------------------------------------------------------------


def test_r1_xlsx_merged_first_row_value_keeps_date_semantics(tmp_path: Path) -> None:
    import datetime as dt

    from openpyxl import Workbook

    from app.shared.core.data_source.loaders.excel_loader import read_merged_ranges_from_xlsx

    wb = Workbook()
    ws = wb.active
    ws.title = "S"
    ws["A1"] = "name"
    ws["B1"] = "deadline"
    ws["A2"], ws["A3"], ws["A4"] = "甲", "乙", "丙"
    ws["B2"] = dt.date(2024, 1, 15)  # openpyxl 对 date 自动应用 yyyy-mm-dd 数字格式
    ws.merge_cells("B2:B4")
    f = tmp_path / "merged.xlsx"
    wb.save(f)

    regions = read_merged_ranges_from_xlsx(f, "S")
    assert regions, "合并区域应被解析到"
    values = regions[0][4]  # (min_row,min_col,max_row,max_col,首行值列表)
    assert isinstance(values[0], (dt.date, dt.datetime)), f"日期语义应保留，当前得到序列号: {values[0]!r}"


# ---------------------------------------------------------------------------
# R6（红线#6）：ai generate --apply 不应清空既有 manifest 的 project.id/settings
# 当前坐实：保留清单仅 transforms/manual_data，LLM 返回空 manifest 时
# project.id 被写空、settings 被整体清除。
# ---------------------------------------------------------------------------


def test_r6_apply_generated_preserves_project_id_and_settings(tmp_path: Path) -> None:
    import yaml

    from app.cli.shared_services.generation_ops import apply_generated_config

    (tmp_path / "project.precis.yaml").write_text(
        "version: 2\n"
        "project:\n  id: my-proj\n  name: 我的项目\n"
        "settings:\n  validation:\n    timeout_seconds: 60\n"
        "schemas: []\nconstraints: []\n",
        encoding="utf-8",
    )

    apply_generated_config(
        {"schemas": {}, "constraints": {}, "regex_nodes": {}},  # LLM 未返回 manifest
        str(tmp_path),
    )

    saved = yaml.safe_load((tmp_path / "project.precis.yaml").read_text(encoding="utf-8"))
    assert saved["project"]["id"] == "my-proj", f"project.id 被清空: {saved.get('project')}"
    assert saved.get("settings", {}).get("validation", {}).get("timeout_seconds") == 60


# ---------------------------------------------------------------------------
# H10（高置信#10）：LLM 生成的实体 id 含路径成分时不应越项目写盘
# 当前坐实：sid='../evil' 拼出 schemas/../evil.schema.yaml 落在项目外，
# write_yaml_atomic(parents=True) 还会建目录——越界写原语零闸门。
# ---------------------------------------------------------------------------


def test_h10_apply_generated_rejects_path_like_id(tmp_path: Path) -> None:
    from app.cli.shared_services.generation_ops import apply_generated_config

    apply_generated_config(
        {
            "manifest": {"version": 2, "project": {"id": "p", "name": "p"}},
            "schemas": {"../../evil": {"version": 2, "name": "evil", "columns": []}},
            "constraints": {},
            "regex_nodes": {},
        },
        str(tmp_path),
    )

    # schemas/../../evil.schema.yaml 经 OS 解析落在项目父目录（越界一层；
    # 单层 ../ 只回到项目根，不构成越界——代理原报告层数有误，此处已修正）
    assert not (tmp_path.parent / "evil.schema.yaml").exists(), "id 含 .. 不应写出项目根之外"


# ---------------------------------------------------------------------------
# R7（红线#7）：单发模式交互命令 Ctrl+C 应友好退出（exit 0/1/2），不应裸 traceback
# 当前坐实：main 单发分支无 KeyboardInterrupt 兜底（仅 REPL 循环有），
# `precis provider` 菜单中 Ctrl+C → traceback + exit 130。
# ---------------------------------------------------------------------------


def test_r7_single_shot_ctrl_c_exits_cleanly(monkeypatch: pytest.MonkeyPatch) -> None:
    import readchar

    from app.cli.shell import main as shell_main

    def _ctrl_c() -> str:
        raise KeyboardInterrupt

    monkeypatch.setattr(readchar, "readkey", _ctrl_c)
    monkeypatch.setattr("app.cli.shell.interactive_menu.readchar.readkey", _ctrl_c)

    try:
        rc = shell_main.main(["provider"])  # 当前 KeyboardInterrupt 穿透单发分支
    except KeyboardInterrupt:
        raise AssertionError("单发模式 Ctrl+C 不应裸穿透（应友好退出 exit 0/1/2）") from None
    assert isinstance(rc, int) and rc in (0, 1, 2), f"应友好退出，实际返回 {rc!r}"
