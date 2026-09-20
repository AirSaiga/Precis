# SPDX-License-Identifier: Apache-2.0
#
# Copyright 2026 Precis Team
#
# Licensed under the Apache License, Version 2.0 (the "License");
# you may not use this file except in compliance with the License.
# You may obtain a copy of the License at
#
#    http://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing, software
# distributed under the License is distributed on an "AS IS" BASIS,
# WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
# See the License for the specific language governing permissions and
# limitations under the License.
"""@fileoverview 校验报告导出测试（P2-3 --report）

覆盖：HTML/Excel 文件生成与关键内容、按扩展名分派、CLI --report 集成
（JSON 模式 stdout 纯净 + 报告内容与 JSON 同源）、不支持的扩展名报错。
"""

from __future__ import annotations

import json

import pytest

from app.shared.services.validation.report_export import export_excel_report, export_html_report, export_report

# 最小契约 payload（与 _build_json_payload 输出同构）
_PAYLOAD = {
    "schema_version": 1,
    "is_valid": False,
    "interrupted": False,
    "duration_ms": 12,
    "tables": [{"name": "orders", "rows": 3}],
    "summary": {"constraints_total": 2, "constraints_passed": 1, "constraints_failed": 1},
    "errors": [
        {
            "table": "orders",
            "column": "amount",
            "constraint_type": "NotNullConstraint",
            "constraint_file": "constraints/c1.constraint.yaml",
            "row_index": 1,
            "cell_value": None,
            "error_message": "非空约束冲突: 列 'amount' 的值不能为空。",
        },
        {
            "table": "orders",
            "column": "unit_price",
            "constraint_type": "RangeConstraint",
            "constraint_file": "constraints/c2.constraint.yaml",
            "row_index": 2,
            "cell_value": 999999.0,
            "error_message": "区间约束冲突: 值 999999.0 不在范围 [0, 100000] 内。",
        },
    ],
    "loading_warnings": [],
}


class TestHtmlReport:
    def test_file_created_with_key_content(self, tmp_path):
        path = export_html_report(_PAYLOAD, tmp_path / "report.html")
        assert path.exists()
        content = path.read_text(encoding="utf-8")
        # 结论、概览、明细关键字段均出现在报告中
        assert "发现违规" in content
        assert "数据校验报告" in content
        assert "非空约束冲突" in content
        assert "constraints/c1.constraint.yaml" in content
        assert "orders" in content
        # 自包含：内联样式、无外部资源引用
        assert "<style>" in content
        assert "http://" not in content and "https://" not in content

    def test_html_escaping(self, tmp_path):
        """单元格值经 HTML 转义，防注入。"""
        payload = dict(_PAYLOAD)
        payload["errors"] = [dict(_PAYLOAD["errors"][0], cell_value="<script>alert(1)</script>")]
        path = export_html_report(payload, tmp_path / "r.html")
        content = path.read_text(encoding="utf-8")
        assert "<script>" not in content
        assert "&lt;script&gt;" in content

    def test_pass_report_shows_no_violations(self, tmp_path):
        payload = dict(_PAYLOAD, is_valid=True, errors=[])
        path = export_html_report(payload, tmp_path / "r.html")
        content = path.read_text(encoding="utf-8")
        assert "通过" in content
        assert "无违规记录" in content


class TestExcelReport:
    def test_file_created_with_sheets_and_rows(self, tmp_path):
        pytest.importorskip("openpyxl")
        from openpyxl import load_workbook

        path = export_excel_report(_PAYLOAD, tmp_path / "report.xlsx")
        assert path.exists()

        workbook = load_workbook(path)
        assert "概览" in workbook.sheetnames
        assert "orders" in workbook.sheetnames
        orders_sheet = workbook["orders"]
        # 表头 + 2 条违规
        assert orders_sheet.max_row == 3
        headers = [c.value for c in orders_sheet[1]]
        assert headers == ["表", "列", "行号", "值", "约束类型", "约束文件", "错误消息"]
        # 违规行第 2 列是列名 amount
        assert orders_sheet.cell(row=2, column=2).value == "amount"

    def test_parent_dir_created(self, tmp_path):
        """输出路径的父目录不存在时自动创建。"""
        target = tmp_path / "nested" / "dir" / "r.xlsx"
        path = export_excel_report(_PAYLOAD, target)
        assert path.exists()


class TestExportDispatch:
    def test_dispatch_by_extension(self, tmp_path):
        html_path = export_report(_PAYLOAD, tmp_path / "r.html")
        xlsx_path = export_report(_PAYLOAD, tmp_path / "r.xlsx")
        assert html_path.suffix == ".html" and html_path.exists()
        assert xlsx_path.suffix == ".xlsx" and xlsx_path.exists()

    def test_unsupported_extension_raises(self, tmp_path):
        with pytest.raises(ValueError, match="不支持的报告格式"):
            export_report(_PAYLOAD, tmp_path / "r.pdf")


class TestCliReportIntegration:
    """CLI --report 与 --format json 并用：stdout 纯 JSON + 报告内容同源"""

    def _make_project(self, tmp_path):
        proj = tmp_path / "proj"
        (proj / "schemas").mkdir(parents=True)
        (proj / "constraints").mkdir()
        (proj / "data").mkdir()
        (proj / "data" / "orders.csv").write_text("id,amount\n1,10\n2,\n", encoding="utf-8")
        (proj / "schemas" / "orders.schema.yaml").write_text(
            """version: 2
id: orders
name: orders
source:
  mode: relative_file
  path: data/orders.csv
columns:
  - id: id
    name: id
    type: integer
  - id: amount
    name: amount
    type: integer
""",
            encoding="utf-8",
        )
        (proj / "constraints" / "c1.constraint.yaml").write_text(
            """version: 2
id: c1a2b3c4-0000-4000-8000-000000000001
type: NotNull
enabled: true
refs:
  table_id: orders
  column_id: amount
params: {}
""",
            encoding="utf-8",
        )
        (proj / "project.precis.yaml").write_text(
            """version: 2
project:
  id: report-demo
  name: report-demo
schemas:
  - id: orders
    path: schemas/orders.schema.yaml
constraints:
  - id: c1a2b3c4-0000-4000-8000-000000000001
    path: constraints/c1.constraint.yaml
""",
            encoding="utf-8",
        )
        return proj

    def test_json_stdout_clean_and_report_written(self, tmp_path, capsys):
        from app.cli.shell.commands.base import ProjectContext
        from app.cli.shell.commands.validate import ValidateCommand

        proj = self._make_project(tmp_path)
        report = tmp_path / "out" / "report.html"
        cmd = ValidateCommand()
        result = cmd.execute(
            ["--manifest", str(proj / "project.precis.yaml"), "--format", "json", "--report", str(report)],
            ProjectContext(),
        )
        captured = capsys.readouterr()

        assert result.success is False
        # stdout 仍是纯 JSON（报告提示走 stderr）
        payload = json.loads(captured.out)
        assert payload["is_valid"] is False
        assert "报告已写入" in captured.err
        # 报告文件存在且内容与 JSON 同源（消息经 HTML 转义后出现）
        assert report.exists()
        content = report.read_text(encoding="utf-8")
        import html as _html

        assert _html.escape(payload["errors"][0]["error_message"]) in content

    def test_unsupported_report_extension_exit_2(self, tmp_path, capsys):
        from app.cli.shell.commands.base import ProjectContext
        from app.cli.shell.commands.validate import ValidateCommand

        proj = self._make_project(tmp_path)
        cmd = ValidateCommand()
        result = cmd.execute(
            ["--manifest", str(proj / "project.precis.yaml"), "--format", "json", "--report", str(tmp_path / "r.pdf")],
            ProjectContext(),
        )
        assert result.success is False
        assert result.exit_code == 2
        assert "报告导出失败" in result.message
