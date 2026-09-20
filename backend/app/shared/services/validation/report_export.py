# SPDX-License-Identifier: Apache-2.0
#
# Copyright 2026 Precis Team
#
# Licensed under the Apache License, Version 2.0 (the "License");
# you may not use this file except in compliance with the License.
# You may obtain a copy of the License at
#
#    http://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing, software
# distributed under the License is distributed on an "AS IS" BASIS,
# WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
# See the License for the specific language governing permissions and
# limitations under the License.
"""@fileoverview 校验报告导出模块（P2-3 --report）

功能概述:
- 把 validate 的 JSON 契约 payload 导出为人类可分享的报告文件
- HTML：自包含单文件（内联 CSS，无外部资源），违规明细表
- Excel：openpyxl（backend 现有依赖），每张表一个 sheet，违规行标红

设计说明:
- 输入直接复用 `_build_json_payload` 的契约结构，保证
  "stdout JSON / 报告文件" 两产物内容一致（验收口径）
- 纯函数 + 显式路径参数，无隐藏 I/O；HTML 值经 html.escape 防注入
"""

from __future__ import annotations

import html
from pathlib import Path
from typing import Any

# 违规明细表的列头（HTML 与 Excel 共用口径）
_REPORT_COLUMNS = ("表", "列", "行号", "值", "约束类型", "约束文件", "错误消息")


def _error_row(entry: dict[str, Any]) -> list[str]:
    """把契约错误条目转为报告行（None 统一展示为空串）。"""
    message = entry.get("error_message") or ""
    suggestion = entry.get("suggestion")
    if suggestion:
        # 修复建议拼进消息单元格，避免两种导出格式各加一列
        message = f"{message}（建议：{suggestion}）"
    return [
        entry.get("table") or "",
        entry.get("column") or "",
        "" if entry.get("row_index") is None else str(entry["row_index"]),
        "" if entry.get("cell_value") is None else str(entry["cell_value"]),
        entry.get("constraint_type") or "",
        entry.get("constraint_file") or "",
        message,
    ]


def _summary_line(payload: dict[str, Any]) -> str:
    """一句话概览：结论 + 检查数 + 耗时。"""
    summary = payload.get("summary", {})
    verdict = "通过" if payload.get("is_valid") else "发现违规"
    return (
        f"校验{verdict}：约束检查 {summary.get('constraints_total', 0)} 项，"
        f"通过 {summary.get('constraints_passed', 0)} / 失败 {summary.get('constraints_failed', 0)}，"
        f"耗时 {payload.get('duration_ms', 0)} ms"
    )


def export_html_report(payload: dict[str, Any], output_path: str | Path) -> Path:
    """导出自包含 HTML 报告。

    Args:
        payload: validate JSON 契约 payload（_build_json_payload 输出）
        output_path: 报告文件路径（.html）

    Returns:
        写入的文件路径

    Raises:
        OSError: 写文件失败
    """
    rows = payload.get("errors", [])
    table_headers = "".join(f"<th>{html.escape(col)}</th>" for col in _REPORT_COLUMNS)
    table_rows: list[str] = []
    for index, entry in enumerate(rows, 1):
        cells = "".join(f"<td>{html.escape(cell)}</td>" for cell in _error_row(entry))
        row_class = "fail" if not payload.get("is_valid") else ""
        table_rows.append(f'<tr class="{row_class}"><td class="idx">{index}</td>{cells}</tr>')
    table_html = "".join(table_rows) if table_rows else '<tr><td colspan="8" class="empty">无违规记录</td></tr>'

    tables_info = (
        "、".join(
            f"{t.get('name') or '?'}({t.get('rows') if t.get('rows') is not None else '?'} 行)"
            for t in payload.get("tables", [])
        )
        or "无"
    )

    document = f"""<!DOCTYPE html>
<html lang="zh-CN">
<head>
<meta charset="utf-8">
<title>Precis 校验报告</title>
<style>
  body {{ font-family: "Segoe UI", "Microsoft YaHei", sans-serif; margin: 2rem; color: #1f2937; }}
  h1 {{ font-size: 1.4rem; }}
  .meta {{ color: #6b7280; margin-bottom: 1rem; }}
  .verdict {{ display: inline-block; padding: 0.2rem 0.8rem; border-radius: 999px; font-weight: 600; }}
  .verdict.pass {{ background: #d1fae5; color: #065f46; }}
  .verdict.fail {{ background: #fee2e2; color: #991b1b; }}
  table {{ border-collapse: collapse; width: 100%; font-size: 0.9rem; }}
  th, td {{ border: 1px solid #d1d5db; padding: 0.4rem 0.6rem; text-align: left; vertical-align: top; }}
  th {{ background: #f3f4f6; }}
  tr.fail td {{ background: #fef2f2; }}
  td.idx {{ color: #9ca3af; width: 2.5rem; }}
  td.empty {{ text-align: center; color: #6b7280; padding: 2rem; }}
</style>
</head>
<body>
<h1>Precis 数据校验报告</h1>
<p class="meta">
  <span class="verdict {"pass" if payload.get("is_valid") else "fail"}">{"通过" if payload.get("is_valid") else "发现违规"}</span>
  &nbsp;{_summary_line(payload)}
  <br>数据表：{html.escape(tables_info)}
</p>
<table>
<thead><tr><th>#</th>{table_headers}</tr></thead>
<tbody>{table_html}</tbody>
</table>
</body>
</html>
"""
    path = Path(output_path)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(document, encoding="utf-8")
    return path


def export_excel_report(payload: dict[str, Any], output_path: str | Path) -> Path:
    """导出 Excel 报告（每张表一个 sheet，违规行标红）。

    Sheet 组织：每张数据表一个 sheet（列出该表违规明细）；无表归属的错误
    进入"汇总" sheet；另有"概览" sheet 记录结论与统计。

    Args:
        payload: validate JSON 契约 payload（_build_json_payload 输出）
        output_path: 报告文件路径（.xlsx）

    Returns:
        写入的文件路径

    Raises:
        OSError: 写文件失败
    """
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    red_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
    header_font = Font(bold=True)

    workbook = Workbook()
    # 概览 sheet
    overview = workbook.active
    overview.title = "概览"
    overview.append(["结论", "通过" if payload.get("is_valid") else "发现违规"])
    for key, label in (
        ("constraints_total", "约束检查总数"),
        ("constraints_passed", "通过数"),
        ("constraints_failed", "失败数"),
    ):
        overview.append([label, payload.get("summary", {}).get(key, 0)])
    overview.append(["耗时(ms)", payload.get("duration_ms", 0)])
    overview.append(["违规条数", len(payload.get("errors", []))])
    overview.append(["数据表", ", ".join(str(t.get("name")) for t in payload.get("tables", []))])
    for cell in overview[1]:
        cell.font = header_font

    # 按表分组违规条目（table 为空/None 的归入"汇总" sheet）
    errors_by_table: dict[str, list[dict[str, Any]]] = {}
    for entry in payload.get("errors", []):
        key = entry.get("table") or "汇总"
        errors_by_table.setdefault(key, []).append(entry)

    sheet_names = {t.get("name") for t in payload.get("tables", []) if t.get("name")}
    for name in sorted(sheet_names | set(errors_by_table.keys())):
        # Excel sheet 名长度上限 31 字符，超长截断
        sheet = workbook.create_sheet(title=name[:31])
        sheet.append(list(_REPORT_COLUMNS))
        for cell in sheet[1]:
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center")
        for entry in errors_by_table.get(name, []):
            sheet.append(_error_row(entry))
            for cell in sheet[sheet.max_row]:
                cell.fill = red_fill
        if not errors_by_table.get(name):
            sheet.append(["（该表无违规）"])

    path = Path(output_path)
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    return path


def export_report(payload: dict[str, Any], output_path: str | Path) -> Path:
    """按扩展名分派报告导出（.html/.htm → HTML；.xlsx → Excel）。

    Args:
        payload: validate JSON 契约 payload
        output_path: 报告文件路径

    Returns:
        写入的文件路径

    Raises:
        ValueError: 扩展名不支持
    """
    suffix = Path(output_path).suffix.lower()
    if suffix in (".html", ".htm"):
        return export_html_report(payload, output_path)
    if suffix == ".xlsx":
        return export_excel_report(payload, output_path)
    raise ValueError(f"不支持的报告格式: {suffix}（支持 .html / .xlsx）")
